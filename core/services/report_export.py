"""CSV/XLSX/PDF report exports with an immutable export audit trail (F-25/F-30)."""

import csv
import re
from datetime import date, datetime
from io import BytesIO
from xml.sax.saxutils import escape

from django.http import HttpResponse
from django.utils import timezone

from core.models import AuditEvent, DataExportLog

_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _safe_csv_value(value):
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(_FORMULA_PREFIXES):
        return f"'{value}"
    return value


def _xlsx_value(value):
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    if value is None:
        return ""
    if isinstance(value, str):
        # Prevent spreadsheet formula injection from student-entered text.
        return _safe_csv_value(value)
    if isinstance(value, (int, float, bool)):
        return value
    return str(value)


def _sections(report):
    """Normalize KPI and detailed reports into titled sections of rows."""
    result = []
    kpi_tiles = report.get("kpi_tiles") or ()
    if kpi_tiles:
        result.append(
            (
                "Key metrics",
                ("Metric", "Value"),
                [(label, value) for label, value, _tone in kpi_tiles],
            )
        )
    if report.get("headers") and report.get("rows") is not None:
        result.append(
            (
                report.get("title") or "Report",
                tuple(report.get("headers") or ()),
                list(report.get("rows") or ()),
            )
        )
    for section in report.get("sections") or ():
        result.append(
            (
                section.get("title") or "Section",
                tuple(section.get("headers") or ()),
                list(section.get("rows") or ()),
            )
        )
    return result


def _filename(report, file_format):
    base = re.sub(r"[^a-z0-9]+", "-", (report.get("title") or "report").lower())
    base = base.strip("-") or "report"
    return f"{base}.{file_format}"


def _filters_for_log(filters):
    cleaned = {}
    for key, value in (filters or {}).items():
        if value is None or value == "":
            continue
        cleaned[key] = str(value)
    return cleaned


def _pdf_text(value):
    text = "" if value is None else str(value)
    return escape(text).replace("\n", "<br/>")


def _pdf_response(report, sections, *, generated_at, filters):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = BytesIO()
    page_size = landscape(A4)
    document = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=report.get("title") or "Report",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontSize=15,
        leading=18,
        spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#475569"),
    )
    heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontSize=10,
        leading=13,
        spaceBefore=8,
        spaceAfter=3,
    )
    cell_style = ParagraphStyle(
        "ReportCell",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
    )
    header_style = ParagraphStyle(
        "ReportHeader",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    story = [Paragraph(_pdf_text(report.get("title") or "Report"), title_style)]
    story.append(
        Paragraph(
            f"Generated {generated_at.strftime('%Y-%m-%d %H:%M')}",
            meta_style,
        )
    )
    if filters:
        story.append(
            Paragraph(
                " • ".join(
                    f"{_pdf_text(key.replace('_', ' ').title())}: {_pdf_text(value)}"
                    for key, value in filters.items()
                ),
                meta_style,
            )
        )
    story.append(Spacer(1, 5 * mm))

    available_width = page_size[0] - 24 * mm
    for section_title, headers, rows in sections:
        story.append(Paragraph(_pdf_text(section_title), heading_style))
        table_data = []
        if headers:
            table_data.append([Paragraph(_pdf_text(value), header_style) for value in headers])
        for row in rows:
            table_data.append([Paragraph(_pdf_text(value), cell_style) for value in row])
        if not table_data:
            story.append(Paragraph("No records for this section.", meta_style))
            continue
        column_count = max(len(row) for row in table_data)
        for row in table_data:
            row.extend([Paragraph("", cell_style)] * (column_count - len(row)))
        column_width = available_width / max(column_count, 1)
        table = Table(
            table_data,
            colWidths=[column_width] * column_count,
            repeatRows=1 if headers else 0,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 4 * mm))

    document.build(story)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    return response


def export_report_response(
    report,
    *,
    file_format,
    actor,
    filters=None,
    filename=None,
    purpose="",
):
    """Build a CSV/XLSX/PDF response and write both export audit records."""
    file_format = (file_format or "csv").lower()
    if file_format not in ("csv", "xlsx", "pdf"):
        raise ValueError("Choose CSV, XLSX, or PDF as the export format.")

    sections = _sections(report)
    row_count = sum(len(rows) for _title, _headers, rows in sections)
    filters = _filters_for_log(filters)
    filename = filename or _filename(report, file_format)
    generated_at = timezone.localtime()

    if file_format == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        writer = csv.writer(response)
        writer.writerow([_safe_csv_value(report.get("title") or "Report")])
        writer.writerow(["Generated", generated_at.strftime("%Y-%m-%d %H:%M")])
        for key, value in filters.items():
            writer.writerow([_safe_csv_value(key.replace("_", " ").title()), value])
        writer.writerow([])
        for title, headers, rows in sections:
            writer.writerow([_safe_csv_value(title)])
            if headers:
                writer.writerow([_safe_csv_value(value) for value in headers])
            for row in rows:
                writer.writerow([_safe_csv_value(value) for value in row])
            writer.writerow([])
    elif file_format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Report"
        sheet["A1"] = report.get("title") or "Report"
        sheet["A1"].font = Font(bold=True, size=14)
        sheet["A2"] = "Generated"
        sheet["B2"] = generated_at.strftime("%Y-%m-%d %H:%M")
        row_number = 3
        for key, value in filters.items():
            sheet.cell(row=row_number, column=1, value=key.replace("_", " ").title())
            sheet.cell(row=row_number, column=2, value=value)
            row_number += 1
        row_number += 1

        for title, headers, rows in sections:
            sheet.cell(row=row_number, column=1, value=title).font = Font(bold=True)
            sheet.cell(row=row_number, column=1).fill = PatternFill(
                "solid", fgColor="E2E8F0"
            )
            row_number += 1
            if headers:
                for column, value in enumerate(headers, start=1):
                    cell = sheet.cell(row=row_number, column=column, value=value)
                    cell.font = Font(bold=True)
                row_number += 1
            for row in rows:
                for column, value in enumerate(row, start=1):
                    sheet.cell(
                        row=row_number,
                        column=column,
                        value=_xlsx_value(value),
                    )
                row_number += 1
            row_number += 1

        column_count = max(
            (len(headers) for _t, headers, _r in sections),
            default=2,
        )
        for column in range(1, column_count + 1):
            width = max(
                (
                    len(str(row[column - 1])) if len(row) >= column else 0
                    for _title, _headers, rows in sections
                    for row in rows
                ),
                default=10,
            )
            sheet.column_dimensions[get_column_letter(column)].width = min(
                max(width + 2, 12), 50
            )

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        workbook.save(response)
    else:
        response = _pdf_response(
            report,
            sections,
            generated_at=generated_at,
            filters=filters,
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'

    purpose = (purpose or "Operational report download").strip()[:255]
    DataExportLog.objects.create(
        actor=actor if getattr(actor, "is_authenticated", False) else None,
        report_name=(report.get("title") or "Report")[:150],
        file_format=file_format,
        filters=filters,
        row_count=row_count,
        purpose=purpose,
    )
    AuditEvent.record(
        actor=actor,
        action=AuditEvent.Action.EXPORT,
        target_type="Report",
        target_id=filename,
        summary=f"Exported {report.get('title') or 'report'} as {file_format}",
        metadata={"row_count": row_count, "filters": filters},
    )
    return response
