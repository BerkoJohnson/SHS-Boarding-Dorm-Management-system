import csv
import io
import random
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Count, Q, Sum
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from core.models import (
    ClassGroup,
    Dormitory,
    DormitoryCapacityException,
    DormitoryRoom,
    Guardian,
    House,
    Programme,
    Student,
)
from core.services.roster import normalize_roster_house_name

MAX_ROSTER_ROWS = 2000
REQUIRED_COLUMNS = {
    "admission_no",
    "index_number",
    "full_name",
    "gender",
    "programme",
}
COLUMN_ALIASES = {
    "admissionno": "admission_no",
    "admissionnumber": "admission_no",
    "indexno": "index_number",
    "indexnumber": "index_number",
    "fullname": "full_name",
    "studentname": "full_name",
    "name": "full_name",
    "gender": "gender",
    "dateofbirth": "date_of_birth",
    "dob": "date_of_birth",
    "programme": "programme",
    "program": "programme",
    "class": "class_group",
    "classgroup": "class_group",
    "form": "class_group",
    "residency": "residency",
    "residencystatus": "residency",
    "house": "house",
    "boardinghouse": "house",
    "dormitory": "dormitory",
    "dorm": "dormitory",
    "block": "dormitory",
    "room": "room",
    "roomname": "room",
    "status": "status",
    "enrollmentstatus": "status",
    "guardianname": "guardian_name",
    "parentname": "guardian_name",
    "guardianphone": "guardian_phone",
    "guardiancontact": "guardian_phone",
    "parentphone": "guardian_phone",
    "parentcontact": "guardian_phone",
    "guardianaltphone": "guardian_alt_phone",
    "parentaltcontact": "guardian_alt_phone",
}


@dataclass
class StagedStudent:
    student: Student
    guardian_name: str
    guardian_phone: str
    guardian_alt_phone: str


@dataclass
class RosterValidation:
    row_count: int = 0
    headers: list[dict] = field(default_factory=list)
    sample_rows: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    staged: list[StagedStudent] = field(default_factory=list, repr=False)
    guardians_by_phone: dict = field(default_factory=dict, repr=False)

    @property
    def is_valid(self):
        return not self.errors and bool(self.staged)


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized_header(value):
    return re.sub(r"[^a-z0-9]", "", _text(value).lower())


def _row_values(headers, values, sheet_name, row_number, errors):
    values = list(values)
    if not any(_text(value) for value in values):
        return None
    if len(values) > len(headers) and any(
        _text(value) for value in values[len(headers) :]
    ):
        errors.append(
            f"{sheet_name}, row {row_number}: data exists beyond the last header."
        )
    mapped = {}
    for index, raw_header in enumerate(headers):
        canonical = COLUMN_ALIASES.get(_normalized_header(raw_header))
        if not canonical:
            continue
        if canonical in mapped:
            errors.append(
                f"{sheet_name}, row {row_number}: column '{raw_header}' duplicates "
                f"the {canonical.replace('_', ' ')} column."
            )
            continue
        mapped[canonical] = _text(values[index] if index < len(values) else "")
    return mapped


def _read_csv(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("CSV files must use UTF-8 encoding.") from error
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        headers = next(reader)
    except StopIteration:
        raise ValueError("The uploaded CSV file is empty.") from None
    return [("CSV", headers, 1, None)] + [
        ("CSV", headers, row_number, row) for row_number, row in enumerate(reader, 2)
    ]


def _read_xlsx(data):
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError) as error:
        raise ValueError(
            "The uploaded file is not a readable Excel .xlsx workbook."
        ) from error

    rows = []
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, None)
            if not headers or not any(_text(value) for value in headers):
                continue
            rows.append((sheet.title, headers, 1, None))
            for row_number, row in enumerate(iterator, 2):
                rows.append((sheet.title, headers, row_number, row))
                if len(rows) > MAX_ROSTER_ROWS + len(workbook.worksheets):
                    raise ValueError(
                        f"Roster files cannot contain more than {MAX_ROSTER_ROWS} student rows."
                    )
    finally:
        workbook.close()
    if not rows:
        raise ValueError("The Excel workbook has no student rows.")
    return rows


def _reference_index(objects, attribute):
    index = {}
    for obj in objects:
        key = _text(getattr(obj, attribute)).casefold()
        index.setdefault(key, []).append(obj)
    return index


def _reference(value, index, field, location, errors):
    if not value:
        return None
    matches = index.get(value.casefold(), [])
    if len(matches) == 1:
        return matches[0]
    if not matches:
        errors.append(f"{location}: unknown {field} '{value}'.")
    else:
        errors.append(f"{location}: {field} '{value}' is ambiguous.")
    return None


def _parse_date(value, location, errors):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    errors.append(f"{location}: date of birth must be YYYY-MM-DD or DD/MM/YYYY.")
    return None


def _choice_value(value, choices, field, location, errors, default=""):
    if not value:
        return default
    normalized = value.casefold()
    for choice_value, label in choices:
        if normalized in {str(choice_value).casefold(), str(label).casefold()}:
            return choice_value
    errors.append(f"{location}: invalid {field} '{value}'.")
    return default


def _balanced_dormitory(dormitories, occupancy, limits):
    available = [
        dormitory
        for dormitory in dormitories
        if occupancy.get(dormitory.pk, 0) < limits[dormitory.pk]
    ]
    candidates = available or dormitories
    if not candidates:
        return None
    least_occupied = min(occupancy.get(dormitory.pk, 0) for dormitory in candidates)
    tied = [
        dormitory
        for dormitory in candidates
        if occupancy.get(dormitory.pk, 0) == least_occupied
    ]
    return tied[0] if len(tied) == 1 else random.choice(tied)


def validate_student_roster(uploaded_file, progress_callback=None):
    """Validate a complete CSV/XLSX roster without creating records."""
    from tenancy.entitlements import feature_enabled

    if not feature_enabled("roster_import"):
        return RosterValidation(
            errors=["This tenant's subscription does not include roster import."]
        )

    extension = Path(uploaded_file.name).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        return RosterValidation(errors=["Upload a CSV or Excel .xlsx file."])
    if uploaded_file.size > 10 * 1024 * 1024:
        return RosterValidation(errors=["Roster files must be 10 MB or smaller."])

    if progress_callback:
        progress_callback("Reading roster", 0, 1)
    data = uploaded_file.read()
    try:
        raw_rows = _read_csv(data) if extension == ".csv" else _read_xlsx(data)
    except ValueError as error:
        return RosterValidation(errors=[str(error)])
    if progress_callback:
        progress_callback("Reading roster", 1, 1)

    errors = []
    parsed_rows = []
    headers_by_sheet = {}
    sample_rows = []
    data_row_count = 0
    checked_sheets = set()
    for sheet_name, headers, row_number, row in raw_rows:
        if sheet_name not in checked_sheets:
            checked_sheets.add(sheet_name)
            mapped_headers = [
                {
                    "source": _text(header),
                    "field": COLUMN_ALIASES.get(_normalized_header(header)),
                }
                for header in headers
            ]
            headers_by_sheet[sheet_name] = mapped_headers
            available = {
                COLUMN_ALIASES.get(_normalized_header(header)) for header in headers
            }
            missing = REQUIRED_COLUMNS - available
            if missing:
                errors.append(
                    f"{sheet_name}: missing required columns: "
                    f"{', '.join(sorted(name.replace('_', ' ') for name in missing))}."
                )
                continue
        if row is None:
            continue
        mapped = _row_values(headers, row, sheet_name, row_number, errors)
        if mapped is not None:
            data_row_count += 1
            parsed_rows.append((sheet_name, row_number, mapped))
            if len(sample_rows) < 10:
                sample_rows.append(
                    {"sheet": sheet_name, "row": row_number, "values": mapped}
                )
    if data_row_count > MAX_ROSTER_ROWS:
        errors.append(
            f"Roster files cannot contain more than {MAX_ROSTER_ROWS} student rows."
        )
    if not parsed_rows and not errors:
        errors.append("The roster file contains no student rows.")
    if errors:
        return RosterValidation(
            row_count=data_row_count,
            headers=[
                {"sheet": sheet, "columns": columns}
                for sheet, columns in headers_by_sheet.items()
            ],
            sample_rows=sample_rows,
            errors=errors,
        )

    if progress_callback:
        progress_callback("Validating rows", 0, len(parsed_rows))
    programmes = list(Programme.objects.all())
    programme_index = _reference_index(programmes, "name")
    for programme in programmes:
        if programme.code:
            code_matches = programme_index.setdefault(programme.code.casefold(), [])
            if all(match.pk != programme.pk for match in code_matches):
                code_matches.append(programme)
    class_index = _reference_index(ClassGroup.objects.all(), "name")
    house_index = _reference_index(House.objects.all(), "name")
    dormitories = list(Dormitory.objects.select_related("house").all())
    dormitory_occupancy = {
        dormitory_id: occupied
        for dormitory_id, occupied in Student.objects.filter(
            residency=Student.ResidencyStatus.BOARDING
        )
        .exclude(status=Student.EnrollmentStatus.COMPLETED)
        .values("dormitory_id")
        .annotate(occupied=Count("pk"))
        .values_list("dormitory_id", "occupied")
    }
    now = timezone.now()
    exception_capacity = dict(
        DormitoryCapacityException.objects.filter(
            dormitory_id__in=[dormitory.pk for dormitory in dormitories],
            cancelled_at__isnull=True,
            starts_at__lte=now,
            ends_at__gte=now,
        )
        .values("dormitory_id")
        .annotate(additional_capacity=Sum("additional_capacity"))
        .values_list("dormitory_id", "additional_capacity")
    )
    dormitory_limits = {
        dormitory.pk: dormitory.capacity
        + (exception_capacity.get(dormitory.pk) or 0)
        for dormitory in dormitories
    }
    rooms = list(
        DormitoryRoom.objects.select_related("dormitory").annotate(
            occupied=Count(
                "students",
                filter=Q(
                    students__residency=Student.ResidencyStatus.BOARDING,
                    students__status__in=(
                        Student.EnrollmentStatus.ENROLLED,
                        Student.EnrollmentStatus.APPROVED,
                        Student.EnrollmentStatus.IN_PROGRESS,
                    ),
                ),
            )
        )
    )
    dormitory_index = {}
    for dormitory in dormitories:
        dormitory_index.setdefault(
            (dormitory.house_id, dormitory.name.casefold()), []
        ).append(dormitory)
    boarding_dormitories = {}
    for dormitory in dormitories:
        boarding_dormitories.setdefault(
            (dormitory.house_id, dormitory.gender), []
        ).append(dormitory)
    rooms_by_dormitory = {}
    room_index = {}
    room_occupancy = {}
    for room in rooms:
        rooms_by_dormitory.setdefault(room.dormitory_id, []).append(room)
        room_index.setdefault((room.dormitory_id, room.name.casefold()), []).append(
            room
        )
        room_occupancy[room.pk] = room.occupied

    existing_admission_numbers = {
        value.casefold()
        for value in Student.objects.values_list("admission_no", flat=True)
    }
    existing_index_numbers = {
        value.casefold()
        for value in Student.objects.values_list("index_number", flat=True)
    }
    seen_admission_numbers = set()
    seen_index_numbers = set()
    roster_guardian_phones = {
        row.get("guardian_phone", "").strip()
        for _, _, row in parsed_rows
        if row.get("guardian_phone", "").strip()
    }
    guardians_by_phone = {}
    for guardian in Guardian.objects.filter(
        phone_number__in=roster_guardian_phones
    ).only("id", "full_name", "phone_number", "alt_phone_number"):
        guardians_by_phone.setdefault(guardian.phone_number, []).append(guardian)

    staged = []
    staged_guardian_names = {}
    for row_index, (sheet_name, row_number, row) in enumerate(parsed_rows, 1):
        if progress_callback and (
            row_index == 1 or row_index % 25 == 0 or row_index == len(parsed_rows)
        ):
            progress_callback("Validating rows", row_index, len(parsed_rows))
        location = f"{sheet_name}, row {row_number}"
        row_error_count = len(errors)
        admission_no = row.get("admission_no", "").strip()
        index_number = row.get("index_number", "").strip()
        for value, seen, existing, label in (
            (
                admission_no,
                seen_admission_numbers,
                existing_admission_numbers,
                "admission number",
            ),
            (
                index_number,
                seen_index_numbers,
                existing_index_numbers,
                "index number",
            ),
        ):
            normalized_value = value.casefold()
            if value and (normalized_value in seen or normalized_value in existing):
                errors.append(f"{location}: duplicate {label} '{value}'.")
            seen.add(normalized_value)

        gender = _choice_value(
            row.get("gender", ""),
            Student.Gender.choices,
            "gender",
            location,
            errors,
        )
        residency = _choice_value(
            row.get("residency", ""),
            Student.ResidencyStatus.choices,
            "residency",
            location,
            errors,
            Student.ResidencyStatus.BOARDING,
        )
        status = _choice_value(
            row.get("status", ""),
            Student.EnrollmentStatus.choices,
            "enrollment status",
            location,
            errors,
            Student.EnrollmentStatus.ENROLLED,
        )
        programme_value = row.get("programme", "").strip()
        programme = _reference(
            programme_value, programme_index, "programme", location, errors
        )
        class_value = row.get("class_group", "").strip()
        if class_value.casefold() == "unallocated":
            class_value = ""
        class_group = (
            _reference(class_value, class_index, "class", location, errors)
            if class_value
            else None
        )
        house_value = row.get("house", "").strip()
        house = (
            _reference(
                normalize_roster_house_name(house_value),
                house_index,
                "house",
                location,
                errors,
            )
            if house_value
            else None
        )
        dormitory_value = row.get("dormitory", "").strip()
        dormitory = None
        if dormitory_value:
            possible_dormitories = (
                dormitory_index.get((house.pk, dormitory_value.casefold()), [])
                if house
                else []
            )
            if len(possible_dormitories) == 1:
                dormitory = possible_dormitories[0]
            elif not house:
                errors.append(
                    f"{location}: select a house before specifying a dormitory."
                )
            elif not possible_dormitories:
                errors.append(
                    f"{location}: unknown dormitory '{dormitory_value}' in {house.name}."
                )
            else:
                errors.append(
                    f"{location}: dormitory '{dormitory_value}' is ambiguous in {house.name}."
                )
        elif residency == Student.ResidencyStatus.BOARDING and house:
            gender_block = None
            if gender == Student.Gender.MALE:
                gender_block = Dormitory.GenderBlock.MALE
            elif gender == Student.Gender.FEMALE:
                gender_block = Dormitory.GenderBlock.FEMALE
            possible_dormitories = (
                boarding_dormitories.get((house.pk, gender_block), [])
                if gender_block
                else []
            )
            if possible_dormitories:
                dormitory = _balanced_dormitory(
                    possible_dormitories,
                    dormitory_occupancy,
                    dormitory_limits,
                )
                if dormitory is not None:
                    row["dormitory"] = dormitory.name
            else:
                errors.append(
                    f"{location}: no {gender or 'gender-compatible'} dormitory "
                    f"is configured in {house.name}."
                )
            if possible_dormitories and dormitory is None:
                errors.append(
                    f"{location}: no gender-compatible dormitory is available."
                )

        room_value = row.get("room", "").strip()
        room = None
        if room_value:
            possible_rooms = (
                room_index.get((dormitory.pk, room_value.casefold()), [])
                if dormitory
                else []
            )
            if len(possible_rooms) == 1:
                room = possible_rooms[0]
            elif not dormitory:
                errors.append(
                    f"{location}: select a dormitory before specifying a room."
                )
            elif not possible_rooms:
                errors.append(
                    f"{location}: unknown room '{room_value}' in {dormitory.name}."
                )
            else:
                errors.append(
                    f"{location}: room '{room_value}' is ambiguous in {dormitory.name}."
                )
        elif dormitory:
            possible_rooms = rooms_by_dormitory.get(dormitory.pk, [])
            if len(possible_rooms) == 1:
                room = possible_rooms[0]

        if room and residency == Student.ResidencyStatus.BOARDING:
            occupied = room_occupancy.get(room.pk, 0)
            if occupied >= room.capacity:
                errors.append(
                    f"{location}: room '{room.name}' is at its capacity of "
                    f"{room.capacity}."
                )
        if (
            dormitory
            and residency == Student.ResidencyStatus.BOARDING
            and status != Student.EnrollmentStatus.COMPLETED
        ):
            occupied = dormitory_occupancy.get(dormitory.pk, 0)
            limit = dormitory_limits[dormitory.pk]
            if occupied >= limit:
                errors.append(
                    f"{location}: dormitory '{dormitory.name}' is at its capacity "
                    f"of {limit}."
                )

        dob = _parse_date(row.get("date_of_birth"), location, errors)
        guardian_name = row.get("guardian_name", "").strip()
        guardian_phone = row.get("guardian_phone", "").strip()
        guardian_alt_phone = row.get("guardian_alt_phone", "").strip()
        if bool(guardian_name) != bool(guardian_phone):
            errors.append(
                f"{location}: guardian name and primary phone must be provided together."
            )
        if guardian_alt_phone and not guardian_phone:
            errors.append(
                f"{location}: provide the guardian primary phone before an alternative phone."
            )
        guardian = None
        if guardian_phone:
            try:
                candidate = Guardian(
                    full_name=guardian_name,
                    phone_number=guardian_phone,
                    alt_phone_number=guardian_alt_phone or None,
                )
                candidate.full_clean()
            except ValidationError as error:
                errors.extend(
                    f"{location}: guardian {field}: {message}"
                    for field, messages in error.message_dict.items()
                    for message in messages
                )
            existing_guardians = guardians_by_phone.get(guardian_phone, [])
            if existing_guardians:
                matching_guardians = [
                    item
                    for item in existing_guardians
                    if item.full_name.casefold() == guardian_name.casefold()
                ]
                if not matching_guardians:
                    errors.append(
                        f"{location}: guardian phone is already linked to another name."
                    )
                else:
                    guardian = matching_guardians[0]
            else:
                staged_guardian_name = staged_guardian_names.get(guardian_phone)
                if (
                    staged_guardian_name
                    and staged_guardian_name.casefold() != guardian_name.casefold()
                ):
                    errors.append(
                        f"{location}: guardian phone is used with different names in the file."
                    )
                else:
                    staged_guardian_names[guardian_phone] = guardian_name

        student = Student(
            admission_no=admission_no,
            index_number=index_number,
            full_name=row.get("full_name", "").strip(),
            gender=gender,
            date_of_birth=dob,
            programme=programme,
            house=house,
            dormitory=dormitory,
            room=room,
            class_group=class_group,
            guardian=guardian,
            residency=residency,
            status=status,
        )
        try:
            # Uniqueness and placement checks are performed in batches above.
            # Model.full_clean() would repeat database lookups for unique
            # identifiers and capacity checks for every row, making large
            # uploads exceed the web worker timeout.
            student.clean_fields(
                exclude=(
                    "programme",
                    "house",
                    "dormitory",
                    "room",
                    "class_group",
                    "guardian",
                )
            )
        except ValidationError as error:
            errors.extend(
                f"{location}: {field}: {message}"
                for field, messages in error.message_dict.items()
                for message in messages
            )
        if status == Student.EnrollmentStatus.COMPLETED:
            errors.append(
                f"{location}: completed students must be processed through "
                "the audited phase-out workflow."
            )
        if residency == Student.ResidencyStatus.DAY and dormitory:
            errors.append(f"{location}: day students cannot be assigned a dormitory.")
        if residency == Student.ResidencyStatus.DAY and room:
            errors.append(f"{location}: day students cannot be assigned a room.")
        if residency == Student.ResidencyStatus.BOARDING and not house:
            errors.append(f"{location}: boarding students must be assigned a house.")
        if residency == Student.ResidencyStatus.BOARDING and not dormitory:
            errors.append(
                f"{location}: boarding students must be assigned a dormitory."
            )
        if (
            dormitory
            and gender
            and dormitory.gender
            != (
                Dormitory.GenderBlock.MALE
                if gender == Student.Gender.MALE
                else Dormitory.GenderBlock.FEMALE
            )
        ):
            errors.append(
                f"{location}: gender mismatch; {dormitory.name} is not compatible "
                "with this student's gender."
            )
        if (
            class_group
            and class_group.programme_id
            and programme
            and programme.pk != class_group.programme_id
        ):
            errors.append(
                f"{location}: class belongs to a different programme."
            )
        if (
            room
            and residency == Student.ResidencyStatus.BOARDING
            and len(errors) == row_error_count
        ):
            room_occupancy[room.pk] += 1
        if (
            dormitory
            and residency == Student.ResidencyStatus.BOARDING
            and status != Student.EnrollmentStatus.COMPLETED
            and len(errors) == row_error_count
        ):
            dormitory_occupancy[dormitory.pk] = (
                dormitory_occupancy.get(dormitory.pk, 0) + 1
            )
        staged.append(
            StagedStudent(
                student=student,
                guardian_name=guardian_name,
                guardian_phone=guardian_phone,
                guardian_alt_phone=guardian_alt_phone,
            )
        )

    return RosterValidation(
        row_count=data_row_count,
        headers=[
            {"sheet": sheet, "columns": columns}
            for sheet, columns in headers_by_sheet.items()
        ],
        sample_rows=sample_rows,
        errors=errors,
        staged=staged,
        guardians_by_phone=guardians_by_phone,
    )


def commit_student_roster(validation, progress_callback=None):
    """Atomically create students from a successful validation report."""
    if validation.errors:
        return 0, validation.errors
    if not validation.staged:
        return 0, ["The roster file contains no student rows."]

    staged = validation.staged
    guardians_by_phone = validation.guardians_by_phone
    enrolled_to_import = sum(
        item.student.status == Student.EnrollmentStatus.ENROLLED for item in staged
    )
    try:
        with transaction.atomic():
            from tenancy.entitlements import (
                ensure_quota_available,
                quota_guard,
            )

            # Hold the tenant quota lock across the cap check and the inserts so
            # two concurrent imports cannot both pass the check and overshoot.
            with quota_guard("max_students"):
                if enrolled_to_import:
                    try:
                        ensure_quota_available(
                            "max_students",
                            Student.objects.filter(
                                status=Student.EnrollmentStatus.ENROLLED
                            ).count(),
                            enrolled_to_import,
                        )
                    except ValidationError as error:
                        return 0, error.messages

                created_guardians = {}
                for row_index, item in enumerate(staged, 1):
                    if item.guardian_phone:
                        guardian = next(
                            (
                                guardian
                                for guardian in guardians_by_phone.get(
                                    item.guardian_phone, []
                                )
                                if guardian.full_name.casefold()
                                == item.guardian_name.casefold()
                            ),
                            None,
                        )
                        if guardian is None:
                            guardian = created_guardians.get(item.guardian_phone)
                        if guardian is None:
                            guardian = Guardian.objects.create(
                                full_name=item.guardian_name,
                                phone_number=item.guardian_phone,
                                alt_phone_number=item.guardian_alt_phone or None,
                            )
                            created_guardians[item.guardian_phone] = guardian
                        item.student.guardian = guardian
                    item.student.save()
                    if progress_callback:
                        progress_callback(
                            "Importing students", row_index, len(staged)
                        )
    except IntegrityError:
        return 0, [
            "The roster could not be imported because a student identifier was "
            "added or changed during validation. No students were imported; "
            "please upload the file again."
        ]
    except ValidationError as error:
        return 0, [
            "The roster could not be imported because a student placement became "
            f"invalid during saving. No students were imported: {' '.join(error.messages)}"
        ]
    return len(staged), []


def import_student_roster(uploaded_file, progress_callback=None):
    """Validate a complete roster, then create it atomically."""
    validation = validate_student_roster(
        uploaded_file,
        progress_callback=progress_callback,
    )
    if validation.errors:
        return 0, validation.errors
    return commit_student_roster(validation, progress_callback=progress_callback)
