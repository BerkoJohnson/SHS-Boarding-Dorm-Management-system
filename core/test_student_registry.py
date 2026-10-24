import csv
import io
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import (
    SimpleTestCase,
    TestCase,
    TransactionTestCase,
    override_settings,
)
from django.urls import reverse
from openpyxl import Workbook
from PIL import Image

from core.forms import StudentRegistrationForm
from core.models import (
    ClassGroup,
    Dormitory,
    DormitoryRoom,
    Guardian,
    House,
    Programme,
    StaffRole,
    Student,
    UserProfile,
)
from core.services.student_import import (
    commit_student_roster,
    import_student_roster,
    validate_student_roster,
)


class ImportRosterCommandTenantTests(SimpleTestCase):
    @override_settings(TENANCY_ENABLED=True)
    def test_requires_schema_in_tenant_mode(self):
        with self.assertRaisesMessage(
            CommandError,
            "TENANCY_ENABLED is True; pass --schema to select the school.",
        ):
            call_command("import_roster", "roster.csv")

    @override_settings(TENANCY_ENABLED=True)
    @patch("core.management.commands.import_roster.import_student_roster")
    @patch("tenancy.models.Tenant.objects")
    def test_import_runs_inside_the_selected_tenant(
        self,
        tenant_objects,
        import_roster,
    ):
        tenant_context = MagicMock()
        tenant_objects.filter.return_value.first.return_value = tenant_context
        import_roster.return_value = (2, [])

        with tempfile.TemporaryDirectory() as directory:
            roster_path = Path(directory, "roster.csv")
            roster_path.write_text("roster contents", encoding="utf-8")

            output = io.StringIO()
            call_command(
                "import_roster",
                str(roster_path),
                schema="school_one",
                stdout=output,
            )

        tenant_objects.filter.assert_called_once_with(schema_name="school_one")
        tenant_context.__enter__.assert_called_once_with()
        tenant_context.__exit__.assert_called_once()
        import_roster.assert_called_once()
        self.assertIn("2 student record(s) imported", output.getvalue())

    @override_settings(TENANCY_ENABLED=True)
    @patch("tenancy.models.Tenant.objects")
    def test_unknown_tenant_schema_is_rejected(self, tenant_objects):
        tenant_objects.filter.return_value.first.return_value = None

        with self.assertRaisesMessage(
            CommandError,
            "Tenant schema 'missing_school' was not found.",
        ):
            call_command(
                "import_roster",
                "roster.csv",
                schema="missing_school",
            )


class StudentRegistryTests(TestCase):
    def setUp(self):
        self.media_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.media_directory.cleanup)
        self.media_settings = override_settings(
            MEDIA_ROOT=self.media_directory.name,
            PRIVATE_STUDENT_PHOTO_ROOT=self.media_directory.name,
        )
        self.media_settings.enable()
        self.addCleanup(self.media_settings.disable)

        self.user = get_user_model().objects.create_user(
            username="student-admin",
            password="password123",
        )
        UserProfile.objects.update_or_create(
            user=self.user,
            defaults={"role": StaffRole.ADMIN},
        )
        self.client.force_login(self.user)
        self.programme = Programme.objects.create(name="General Science", code="GSCI")
        self.house = House.objects.create(name="Red House")
        self.male_dorm = Dormitory.objects.create(
            house=self.house,
            name="Boys Block",
            gender=Dormitory.GenderBlock.MALE,
            capacity=80,
        )
        self.female_dorm = Dormitory.objects.create(
            house=self.house,
            name="Girls Block",
            gender=Dormitory.GenderBlock.FEMALE,
            capacity=80,
        )
        self.class_group = ClassGroup.objects.create(
            name="1A1",
            programme=self.programme,
        )

    def _csv_upload(self, rows, headers=None, name="roster.csv"):
        headers = headers or [
            "Admission No.",
            "Index Number",
            "Full Name",
            "Gender",
            "Programme",
            "House",
            "Class",
        ]
        content = io.StringIO()
        writer = csv.writer(content)
        writer.writerow(headers)
        writer.writerows(rows)
        return SimpleUploadedFile(
            name,
            content.getvalue().encode("utf-8"),
            content_type="text/csv",
        )

    def _xlsx_upload(self, rows, headers=None):
        headers = headers or [
            "Admission No.",
            "Index Number",
            "Full Name",
            "Gender",
            "Programme",
            "House",
            "Class",
        ]
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Students"
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        content = io.BytesIO()
        workbook.save(content)
        workbook.close()
        return SimpleUploadedFile(
            "roster.xlsx",
            content.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    def test_registry_and_registration_are_admin_only(self):
        self.client.logout()
        housemaster = get_user_model().objects.create_user(
            username="housemaster",
            password="password123",
        )
        UserProfile.objects.update_or_create(
            user=housemaster,
            defaults={"role": StaffRole.HOUSEMASTER},
        )
        self.client.force_login(housemaster)

        for name in (
            "core:student_registry",
            "core:student_create",
            "core:student_registration_options",
            "core:student_import",
            "core:student_import_template",
        ):
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 302, name)
            self.assertEqual(response["Location"], reverse("core:dashboard"))

        self.client.logout()
        response = self.client.get(reverse("core:student_registry"))
        self.assertEqual(response.status_code, 302)

        self.client.force_login(self.user)
        self.assertEqual(
            self.client.get(reverse("core:student_registry")).status_code,
            200,
        )

    def create_boarding_student(
        self,
        admission_no,
        *,
        full_name="Registry Student",
        gender=Student.Gender.MALE,
        dormitory=None,
        room=None,
        status=Student.EnrollmentStatus.ENROLLED,
    ):
        dormitory = dormitory or self.male_dorm
        return Student.objects.create(
            admission_no=admission_no,
            index_number=f"IDX-{admission_no}",
            full_name=full_name,
            gender=gender,
            programme=self.programme,
            class_group=self.class_group,
            house=dormitory.house,
            dormitory=dormitory,
            room=room,
            status=status,
        )

    def test_registry_search_matches_student_and_boarding_details(self):
        room = DormitoryRoom.objects.create(
            dormitory=self.male_dorm,
            name="North Wing 12",
            capacity=2,
        )
        student = self.create_boarding_student(
            "SEARCH-001",
            full_name="Kwame Search",
            room=room,
        )

        for term in ("Kwame Search", "SEARCH-001", "IDX-SEARCH-001", "Boys Block", "North Wing"):
            with self.subTest(term=term):
                response = self.client.get(
                    reverse("core:student_registry"),
                    {"q": term},
                )
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, student.full_name)

    def test_registry_search_requires_two_characters_and_returns_htmx_partial(self):
        student = self.create_boarding_student("HTMX-001", full_name="Htmx Student")

        page_response = self.client.get(reverse("core:student_registry"))
        self.assertContains(page_response, 'showFilters: false')
        self.assertContains(page_response, 'x-show="showFilters"')
        self.assertContains(page_response, ">Filters</span>")

        filtered_response = self.client.get(
            reverse("core:student_registry"),
            {"gender": Student.Gender.MALE},
        )
        self.assertContains(filtered_response, 'showFilters: true')
        self.assertContains(filtered_response, "All genders")

        short_response = self.client.get(
            reverse("core:student_registry"),
            {"q": "H"},
            HTTP_HX_REQUEST="true",
        )
        self.assertContains(short_response, "Type at least 2 characters to search")
        self.assertNotContains(short_response, student.full_name)

        response = self.client.get(
            reverse("core:student_registry"),
            {"q": "Htmx"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, student.full_name)
        self.assertNotContains(response, "Student Registry")

    def test_registry_filters_can_be_combined(self):
        room = DormitoryRoom.objects.create(
            dormitory=self.female_dorm,
            name="Filtered Room",
            capacity=2,
        )
        female_student = self.create_boarding_student(
            "FILTER-001",
            full_name="Filtered Female",
            gender=Student.Gender.FEMALE,
            dormitory=self.female_dorm,
            room=room,
            status=Student.EnrollmentStatus.APPROVED,
        )
        self.create_boarding_student("FILTER-002", full_name="Other Male")

        response = self.client.get(
            reverse("core:student_registry"),
            {
                "gender": Student.Gender.FEMALE,
                "status": Student.EnrollmentStatus.APPROVED,
                "residency": Student.ResidencyStatus.BOARDING,
                "house": str(self.house.pk),
                "dormitory": str(self.female_dorm.pk),
                "room": str(room.pk),
                "programme": str(self.programme.pk),
                "class_group": str(self.class_group.pk),
            },
        )

        self.assertContains(response, female_student.full_name)
        self.assertNotContains(response, "Other Male")

    def test_registry_paginates_and_keeps_search_and_filter_parameters(self):
        students = [
            Student(
                admission_no=f"PAGE-{number:03}",
                index_number=f"IDX-PAGE-{number:03}",
                full_name=f"Paged Student {number:03}",
                gender=Student.Gender.MALE,
                programme=self.programme,
                class_group=self.class_group,
                house=self.house,
                dormitory=self.male_dorm,
            )
            for number in range(51)
        ]
        Student.objects.bulk_create(students)

        response = self.client.get(
            reverse("core:student_registry"),
            {"q": "Paged Student", "residency": Student.ResidencyStatus.BOARDING},
        )

        self.assertContains(response, "of 51")
        self.assertContains(response, "Page 1 of 6")
        self.assertContains(response, "q=Paged+Student")
        self.assertContains(response, "residency=Boarding")
        self.assertEqual(response.context["students_page"].paginator.per_page, 10)

        next_page = self.client.get(
            reverse("core:student_registry"),
            {
                "q": "Paged Student",
                "residency": Student.ResidencyStatus.BOARDING,
                "page": 2,
            },
        )
        self.assertContains(next_page, "Paged Student 010")
        self.assertNotContains(next_page, "Paged Student 009")
        self.assertContains(next_page, "Page 2 of 6")

    def test_admin_student_search_filters_and_pagination(self):
        room = DormitoryRoom.objects.create(
            dormitory=self.male_dorm,
            name="Admin Search Room",
            capacity=2,
        )
        student = self.create_boarding_student(
            "ADMIN-SEARCH-001",
            full_name="Admin Search Student",
            room=room,
        )
        administrator = get_user_model().objects.create_superuser(
            username="registry-admin",
            email="registry-admin@example.test",
            password="test-password",
        )
        self.client.force_login(administrator)

        for term in (
            "Admin Search Student",
            student.admission_no,
            student.index_number,
            "General Science",
            "1A1",
            "Red House",
            "Boys Block",
            "Admin Search Room",
        ):
            with self.subTest(term=term):
                search_response = self.client.get(
                    reverse("admin:core_student_changelist"),
                    {"q": term},
                )
                self.assertContains(search_response, student.full_name)
        self.assertContains(search_response, "Room")
        self.assertContains(search_response, "Programme")

        students = [
            Student(
                admission_no=f"ADMIN-PAGE-{number:03}",
                index_number=f"IDX-ADMIN-PAGE-{number:03}",
                full_name=f"Admin Paged Student {number:03}",
                gender=Student.Gender.MALE,
                programme=self.programme,
                class_group=self.class_group,
                house=self.house,
                dormitory=self.male_dorm,
            )
            for number in range(51)
        ]
        Student.objects.bulk_create(students)
        page_response = self.client.get(reverse("admin:core_student_changelist"))

        changelist = page_response.context["cl"]
        self.assertEqual(changelist.result_count, 52)
        self.assertEqual(changelist.paginator.per_page, 50)
        self.assertEqual(changelist.paginator.num_pages, 2)

    def test_individual_registration_creates_student_guardian_and_photo(self):
        image = Image.new("RGB", (160, 220), color="navy")
        image_data = io.BytesIO()
        image.save(image_data, format="JPEG")
        image_data.seek(0)
        photo = SimpleUploadedFile(
            "portrait.jpg",
            image_data.read(),
            content_type="image/jpeg",
        )
        payload = {
            "admission_no": "ADM-001",
            "index_number": "IDX-001",
            "full_name": "Ama Mensah",
            "gender": Student.Gender.FEMALE,
            "date_of_birth": "2009-04-10",
            "programme": self.programme.pk,
            "class_group": self.class_group.pk,
            "residency": Student.ResidencyStatus.BOARDING,
            "house": self.house.pk,
            "dormitory": self.female_dorm.pk,
            "status": Student.EnrollmentStatus.ENROLLED,
            "guardian_name": "Akosua Mensah",
            "guardian_phone": "+233240000001",
        }

        response = self.client.post(
            reverse("core:student_create"),
            {**payload, "photo": photo},
        )

        self.assertRedirects(response, reverse("core:student_registry"))
        student = Student.objects.get(admission_no="ADM-001")
        self.assertEqual(student.class_group, self.class_group)
        self.assertEqual(student.guardian.full_name, "Akosua Mensah")
        self.assertTrue(student.photo.name.startswith("students/photos/"))
        self.assertEqual(student.photo.name.rsplit(".", 1)[-1], "jpg")
        self.assertEqual(Guardian.objects.count(), 1)

    def test_registration_form_limits_dependent_querysets_to_selected_parents(self):
        other_programme = Programme.objects.create(name="General Arts", code="GART")
        other_class = ClassGroup.objects.create(
            name="1A2",
            programme=other_programme,
        )
        other_house = House.objects.create(name="Blue House")
        other_dormitory = Dormitory.objects.create(
            house=other_house,
            name="Other Boys Block",
            gender=Dormitory.GenderBlock.MALE,
            capacity=80,
        )
        room = DormitoryRoom.objects.create(
            dormitory=self.male_dorm,
            name="North Room",
            capacity=4,
        )
        DormitoryRoom.objects.create(
            dormitory=other_dormitory,
            name="South Room",
            capacity=4,
        )

        form = StudentRegistrationForm(
            initial={
                "programme": self.programme.pk,
                "house": self.house.pk,
                "dormitory": self.male_dorm.pk,
            }
        )

        self.assertEqual(
            list(form.fields["class_group"].queryset),
            [self.class_group],
        )
        self.assertNotIn(other_class, form.fields["class_group"].queryset)
        self.assertEqual(
            list(form.fields["dormitory"].queryset),
            [self.male_dorm, self.female_dorm],
        )
        self.assertEqual(
            list(form.fields["room"].queryset),
            [room],
        )

    def test_registration_options_return_only_matching_child_records(self):
        other_programme = Programme.objects.create(name="General Arts", code="GART")
        other_class = ClassGroup.objects.create(
            name="1A2",
            programme=other_programme,
        )
        other_house = House.objects.create(name="Blue House")
        other_dormitory = Dormitory.objects.create(
            house=other_house,
            name="Other Boys Block",
            gender=Dormitory.GenderBlock.MALE,
            capacity=80,
        )
        room = DormitoryRoom.objects.create(
            dormitory=self.male_dorm,
            name="North Room",
            capacity=4,
        )
        DormitoryRoom.objects.create(
            dormitory=other_dormitory,
            name="South Room",
            capacity=4,
        )
        options_url = reverse("core:student_registration_options")

        class_response = self.client.get(
            options_url,
            {"type": "class_groups", "programme_id": self.programme.pk},
        )
        dormitory_response = self.client.get(
            options_url,
            {
                "type": "dormitories",
                "house_id": self.house.pk,
                "gender": Student.Gender.MALE,
            },
        )
        female_dormitory_response = self.client.get(
            options_url,
            {
                "type": "dormitories",
                "house_id": self.house.pk,
                "gender": Student.Gender.FEMALE,
            },
        )
        room_response = self.client.get(
            options_url,
            {
                "type": "rooms",
                "house_id": self.house.pk,
                "dormitory_id": self.male_dorm.pk,
                "gender": Student.Gender.MALE,
            },
        )
        mismatched_room_response = self.client.get(
            options_url,
            {
                "type": "rooms",
                "house_id": self.house.pk,
                "dormitory_id": other_dormitory.pk,
                "gender": Student.Gender.MALE,
            },
        )
        mismatched_gender_room_response = self.client.get(
            options_url,
            {
                "type": "rooms",
                "house_id": self.house.pk,
                "dormitory_id": self.male_dorm.pk,
                "gender": Student.Gender.FEMALE,
            },
        )
        invalid_gender_response = self.client.get(
            options_url,
            {
                "type": "dormitories",
                "house_id": self.house.pk,
                "gender": "unknown",
            },
        )

        self.assertEqual(
            [item["id"] for item in class_response.json()["results"]],
            [str(self.class_group.pk)],
        )
        self.assertNotIn(
            str(other_class.pk),
            [item["id"] for item in class_response.json()["results"]],
        )
        self.assertEqual(
            [item["id"] for item in dormitory_response.json()["results"]],
            [str(self.male_dorm.pk)],
        )
        self.assertEqual(
            [item["id"] for item in female_dormitory_response.json()["results"]],
            [str(self.female_dorm.pk)],
        )
        self.assertEqual(
            [item["id"] for item in room_response.json()["results"]],
            [str(room.pk)],
        )
        self.assertEqual(mismatched_room_response.status_code, 404)
        self.assertEqual(mismatched_gender_room_response.status_code, 404)
        self.assertEqual(invalid_gender_response.status_code, 400)

    def test_registration_form_filters_dormitories_by_student_gender(self):
        male_student = Student.objects.create(
            admission_no="ADM-MALE-GENDER",
            index_number="IDX-MALE-GENDER",
            full_name="Male Gender Student",
            gender=Student.Gender.MALE,
            programme=self.programme,
            class_group=self.class_group,
            house=self.house,
            dormitory=self.male_dorm,
        )

        edit_form = StudentRegistrationForm(instance=male_student)
        create_form = StudentRegistrationForm(
            initial={
                "gender": Student.Gender.FEMALE,
                "house": self.house.pk,
            }
        )

        self.assertEqual(
            list(edit_form.fields["dormitory"].queryset),
            [self.male_dorm],
        )
        self.assertEqual(
            list(create_form.fields["dormitory"].queryset),
            [self.female_dorm],
        )

        # A legacy/imported mismatch must stay visible on the edit form so the
        # model validation can report it instead of silently dropping the dorm.
        Student.objects.filter(pk=male_student.pk).update(
            dormitory=self.female_dorm
        )
        male_student.refresh_from_db()
        legacy_form = StudentRegistrationForm(instance=male_student)
        self.assertIn(self.female_dorm, legacy_form.fields["dormitory"].queryset)

    def test_student_gender_must_match_the_dormitory(self):
        with self.assertRaises(ValidationError) as error:
            Student.objects.create(
                admission_no="ADM-GENDER-MISMATCH",
                index_number="IDX-GENDER-MISMATCH",
                full_name="Mismatched Gender Student",
                gender=Student.Gender.MALE,
                programme=self.programme,
                class_group=self.class_group,
                house=self.house,
                dormitory=self.female_dorm,
            )

        self.assertIn(
            "Gender Mismatch",
            error.exception.message_dict["dormitory"][0],
        )

    def test_registration_page_renders_options_without_javascript(self):
        response = self.client.get(reverse("core:student_create"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "JavaScript is turned off.")
        self.assertContains(response, "data-student-registration-options=")
        form = response.context["form"]
        self.assertIn(self.class_group, form.fields["class_group"].queryset)

    def test_new_registration_form_offers_options_without_javascript(self):
        DormitoryRoom.objects.create(
            dormitory=self.male_dorm,
            name="Fallback Room",
            capacity=4,
        )

        form = StudentRegistrationForm()

        self.assertIn(self.class_group, form.fields["class_group"].queryset)
        self.assertIn(self.male_dorm, form.fields["dormitory"].queryset)
        self.assertIn(self.female_dorm, form.fields["dormitory"].queryset)
        self.assertTrue(form.fields["room"].queryset.exists())

    def test_registration_form_keeps_a_class_group_whose_programme_is_unset(self):
        legacy_class = ClassGroup.objects.create(name="LEG1", programme=None)
        student = Student.objects.create(
            admission_no="ADM-LEGACY",
            index_number="IDX-LEGACY",
            full_name="Legacy Class Student",
            gender=Student.Gender.MALE,
            programme=self.programme,
            class_group=legacy_class,
            house=self.house,
            dormitory=self.male_dorm,
        )

        form = StudentRegistrationForm(instance=student)

        self.assertIn(legacy_class, form.fields["class_group"].queryset)

    def test_registration_form_edit_preserves_a_class_group_whose_programme_is_unset(self):
        legacy_class = ClassGroup.objects.create(name="LEG2", programme=None)
        student = Student.objects.create(
            admission_no="ADM-LEGACY-EDIT",
            index_number="IDX-LEGACY-EDIT",
            full_name="Legacy Edit Student",
            gender=Student.Gender.MALE,
            programme=self.programme,
            class_group=legacy_class,
            house=self.house,
            dormitory=self.male_dorm,
        )

        form = StudentRegistrationForm(
            data={
                "admission_no": student.admission_no,
                "index_number": student.index_number,
                "full_name": student.full_name,
                "gender": student.gender,
                "programme": self.programme.pk,
                "class_group": legacy_class.pk,
                "residency": student.residency,
                "house": self.house.pk,
                "dormitory": self.male_dorm.pk,
                "status": student.status,
            },
            instance=student,
        )

        self.assertTrue(form.is_valid(), form.errors)
        saved = form.save()
        saved.refresh_from_db()
        self.assertEqual(saved.class_group, legacy_class)

    def test_registration_options_include_classes_with_no_programme(self):
        legacy_class = ClassGroup.objects.create(name="LEG3", programme=None)

        response = self.client.get(
            reverse("core:student_registration_options"),
            {"type": "class_groups", "programme_id": self.programme.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            str(legacy_class.pk),
            [item["id"] for item in response.json()["results"]],
        )

    def test_registration_form_updates_existing_guardian_alternative_phone(self):
        guardian = Guardian.objects.create(
            full_name="Akosua Mensah",
            phone_number="+233240000001",
            alt_phone_number="+233240000002",
        )
        student = Student.objects.create(
            admission_no="ADM-GUARDIAN",
            index_number="IDX-GUARDIAN",
            full_name="Ama Mensah",
            gender=Student.Gender.FEMALE,
            programme=self.programme,
            house=self.house,
            dormitory=self.female_dorm,
            class_group=self.class_group,
            guardian=guardian,
        )
        form = StudentRegistrationForm(
            data={
                "admission_no": student.admission_no,
                "index_number": student.index_number,
                "full_name": student.full_name,
                "gender": student.gender,
                "programme": self.programme.pk,
                "class_group": self.class_group.pk,
                "residency": student.residency,
                "house": self.house.pk,
                "dormitory": self.female_dorm.pk,
                "status": student.status,
                "guardian_name": guardian.full_name,
                "guardian_phone": guardian.phone_number,
                "guardian_alt_phone": "+233240000003",
            },
            instance=student,
        )

        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        guardian.refresh_from_db()
        self.assertEqual(guardian.alt_phone_number, "+233240000003")

    def test_photo_upload_rejects_fake_image_content(self):
        invalid_photo = SimpleUploadedFile(
            "portrait.jpg",
            b"not an actual image",
            content_type="image/jpeg",
        )
        form = StudentRegistrationForm(
            data={
                "admission_no": "ADM-BAD",
                "index_number": "IDX-BAD",
                "full_name": "Bad Image",
                "gender": Student.Gender.MALE,
                "programme": self.programme.pk,
                "residency": Student.ResidencyStatus.BOARDING,
                "house": self.house.pk,
                "dormitory": self.male_dorm.pk,
                "status": Student.EnrollmentStatus.ENROLLED,
            },
            files={"photo": invalid_photo},
        )

        self.assertFalse(form.is_valid())
        self.assertIn("photo", form.errors)

    def test_csv_import_validates_then_creates_with_configured_placement(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-CSV-1",
                    "IDX-CSV-1",
                    "Kwame Boateng",
                    "Male",
                    "GSCI",
                    "Red House",
                    "1A1",
                ],
                [
                    "ADM-CSV-2",
                    "IDX-CSV-2",
                    "Akua Boateng",
                    "Female",
                    "General Science",
                    "Red House",
                    "1A1",
                ],
            ]
        )

        imported, errors = import_student_roster(uploaded_file)

        self.assertEqual(errors, [])
        self.assertEqual(imported, 2)
        male_student = Student.objects.get(admission_no="ADM-CSV-1")
        female_student = Student.objects.get(admission_no="ADM-CSV-2")
        self.assertEqual(male_student.dormitory, self.male_dorm)
        self.assertEqual(female_student.dormitory, self.female_dorm)
        self.assertEqual(male_student.residency, Student.ResidencyStatus.BOARDING)

    def test_validation_preview_uses_import_validation_without_creating_records(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-PREVIEW",
                    "IDX-PREVIEW",
                    "Preview Student",
                    "Male",
                    "GSCI",
                    "Red House",
                    "1A1",
                ],
            ],
            headers=[
                "Admission No.",
                "Index Number",
                "Full Name",
                "Gender",
                "Programme",
                "House",
                "Class",
                "Extra Notes",
            ],
        )

        validation = validate_student_roster(uploaded_file)

        self.assertTrue(validation.is_valid)
        self.assertEqual(validation.row_count, 1)
        self.assertEqual(validation.headers[0]["sheet"], "CSV")
        self.assertIsNone(validation.headers[0]["columns"][-1]["field"])
        self.assertEqual(
            validation.sample_rows[0]["values"]["full_name"],
            "Preview Student",
        )
        self.assertFalse(Student.objects.exists())

        imported, errors = commit_student_roster(validation)

        self.assertEqual((imported, errors), (1, []))
        self.assertTrue(Student.objects.filter(admission_no="ADM-PREVIEW").exists())

    def test_roster_validation_avoids_per_student_full_clean_queries(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-NO-FULL-CLEAN",
                    "IDX-NO-FULL-CLEAN",
                    "Validation Student",
                    "Male",
                    "GSCI",
                    "Red House",
                    "1A1",
                ],
            ]
        )

        with patch.object(
            Student,
            "full_clean",
            side_effect=AssertionError("per-row full_clean must not run"),
        ):
            validation = validate_student_roster(uploaded_file)

        self.assertTrue(validation.is_valid, validation.errors)

    def test_roster_validation_checks_day_student_placement_without_full_clean(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-DAY-DORM",
                    "IDX-DAY-DORM",
                    "Day Student",
                    "Male",
                    "GSCI",
                    "Red House",
                    "1A1",
                    "Day",
                    "Boys Block",
                ],
            ],
            headers=[
                "Admission No.",
                "Index Number",
                "Full Name",
                "Gender",
                "Programme",
                "House",
                "Class",
                "Residency",
                "Dormitory",
            ],
        )

        validation = validate_student_roster(uploaded_file)

        self.assertFalse(validation.is_valid)
        self.assertTrue(
            any(
                "day students cannot be assigned a dormitory" in error.lower()
                for error in validation.errors
            ),
            validation.errors,
        )

    def test_web_preview_and_confirmation_share_the_service_validation(self):
        row = [
            "ADM-WEB-PREVIEW",
            "IDX-WEB-PREVIEW",
            "Web Preview Student",
            "Male",
            "GSCI",
            "Red House",
            "1A1",
        ]
        preview_response = self.client.post(
            reverse("core:student_import"),
            {"intent": "preview", "roster_file": self._csv_upload([row])},
            HTTP_ACCEPT="application/json",
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(preview_response.status_code, 200)
        preview = preview_response.json()
        self.assertTrue(preview["valid"])
        self.assertEqual(preview["row_count"], 1)
        self.assertFalse(Student.objects.exists())

        import_response = self.client.post(
            reverse("core:student_import"),
            {"intent": "import", "roster_file": self._csv_upload([row])},
            HTTP_ACCEPT="application/json",
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(import_response.status_code, 200)
        self.assertTrue(import_response.json()["success"])
        self.assertTrue(
            Student.objects.filter(admission_no="ADM-WEB-PREVIEW").exists()
        )

    def test_web_preview_renders_column_mapping_and_sample_without_importing(self):
        row = [
            "ADM-HTML-PREVIEW",
            "IDX-HTML-PREVIEW",
            "HTML Preview Student",
            "Male",
            "GSCI",
            "Red House",
            "1A1",
        ]

        response = self.client.post(
            reverse("core:student_import"),
            {"intent": "preview", "roster_file": self._csv_upload([row])},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Roster preview")
        self.assertContains(response, "Admission No.")
        self.assertContains(response, "HTML Preview Student")
        self.assertContains(response, "Confirm and import")
        self.assertFalse(Student.objects.exists())

    def test_web_preview_reports_missing_programme_and_house_before_import(self):
        response = self.client.post(
            reverse("core:student_import"),
            {
                "intent": "preview",
                "roster_file": self._csv_upload(
                    [
                        [
                            "ADM-UNKNOWN-REF",
                            "IDX-UNKNOWN-REF",
                            "Unknown Reference Student",
                            "Male",
                            "New Programme",
                            "New House",
                            "",
                        ]
                    ]
                ),
            },
            HTTP_ACCEPT="application/json",
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 200)
        preview = response.json()
        self.assertFalse(preview["valid"])
        self.assertTrue(any("unknown programme" in error for error in preview["errors"]))
        self.assertTrue(any("unknown house" in error for error in preview["errors"]))
        self.assertFalse(Student.objects.exists())

    def test_cli_dry_run_reports_preview_without_importing(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-DRY-RUN",
                    "IDX-DRY-RUN",
                    "Dry Run Student",
                    "Male",
                    "GSCI",
                    "Red House",
                    "1A1",
                ],
            ]
        )
        with tempfile.TemporaryDirectory() as directory:
            roster_path = Path(directory, uploaded_file.name)
            roster_path.write_bytes(uploaded_file.read())
            output = io.StringIO()

            call_command(
                "import_roster",
                str(roster_path),
                "--dry-run",
                stdout=output,
            )

        self.assertIn("Roster validation passed", output.getvalue())
        self.assertIn("Rows validated: 1", output.getvalue())
        self.assertIn("Dry Run Student", output.getvalue())
        self.assertFalse(Student.objects.exists())

    def test_xlsx_import_normalizes_house_and_treats_unallocated_class_as_blank(self):
        uploaded_file = self._xlsx_upload(
            [
                [
                    "ADM-UNALLOCATED",
                    "IDX-UNALLOCATED",
                    "Unallocated Student",
                    "Male",
                    "GSCI",
                    "Blue House(M), TWO",
                    "Unallocated",
                ],
            ]
        )
        blue_house = House.objects.create(name="House Two (Blue)")
        blue_male_dorm = Dormitory.objects.create(
            house=blue_house,
            name="Boys Block",
            gender=Dormitory.GenderBlock.MALE,
            capacity=80,
        )
        Dormitory.objects.create(
            house=blue_house,
            name="Girls Block",
            gender=Dormitory.GenderBlock.FEMALE,
            capacity=80,
        )

        imported, errors = import_student_roster(uploaded_file)

        self.assertEqual(errors, [])
        self.assertEqual(imported, 1)
        student = Student.objects.get(admission_no="ADM-UNALLOCATED")
        self.assertIsNone(student.class_group)
        self.assertEqual(student.house, blue_house)
        self.assertEqual(student.dormitory, blue_male_dorm)

    def test_excel_import_reads_each_worksheet(self):
        headers = [
            "Admission No.",
            "Index Number",
            "Full Name",
            "Gender",
            "Programme",
            "House",
            "Class",
        ]
        workbook = Workbook()
        first = workbook.active
        first.title = "First"
        first.append(headers)
        first.append(
            [
                "ADM-XLS-1",
                "IDX-XLS-1",
                "First Student",
                "Male",
                "General Science",
                "Red House",
                "1A1",
            ]
        )
        second = workbook.create_sheet("Second")
        second.append(headers)
        second.append(
            [
                "ADM-XLS-2",
                "IDX-XLS-2",
                "Second Student",
                "Female",
                "General Science",
                "Red House",
                "1A1",
            ]
        )
        content = io.BytesIO()
        workbook.save(content)
        workbook.close()
        uploaded_file = SimpleUploadedFile(
            "roster.xlsx",
            content.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        imported, errors = import_student_roster(uploaded_file)

        self.assertEqual(errors, [])
        self.assertEqual(imported, 2)
        self.assertEqual(Student.objects.filter(admission_no__startswith="ADM-XLS").count(), 2)

    def test_invalid_or_duplicate_csv_imports_nothing(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-DUP",
                    "IDX-DUP-1",
                    "First Student",
                    "Male",
                    "General Science",
                    "Red House",
                    "1A1",
                ],
                [
                    "ADM-DUP",
                    "IDX-DUP-2",
                    "Second Student",
                    "Female",
                    "General Science",
                    "Red House",
                    "1A1",
                ],
            ]
        )

        imported, errors = import_student_roster(uploaded_file)

        self.assertEqual(imported, 0)
        self.assertTrue(any("duplicate admission number" in error for error in errors))
        self.assertFalse(Student.objects.exists())

    def test_invalid_placement_or_unknown_references_reject_the_whole_file(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-GOOD",
                    "IDX-GOOD",
                    "Good Student",
                    "Male",
                    "General Science",
                    "Red House",
                    "1A1",
                ],
                [
                    "ADM-BAD",
                    "IDX-BAD",
                    "Bad Student",
                    "Female",
                    "Unknown Programme",
                    "Unknown House",
                    "1A1",
                ],
            ]
        )

        imported, errors = import_student_roster(uploaded_file)

        self.assertEqual(imported, 0)
        self.assertTrue(any("unknown programme" in error for error in errors))
        self.assertTrue(any("unknown house" in error for error in errors))
        self.assertFalse(Student.objects.exists())

    def test_import_rejects_class_from_a_different_programme(self):
        other_programme = Programme.objects.create(
            name="General Arts",
            code="GART",
        )
        ClassGroup.objects.create(name="1A2", programme=other_programme)
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-MISMATCH",
                    "IDX-MISMATCH",
                    "Mismatched Student",
                    "Male",
                    "General Science",
                    "Red House",
                    "1A2",
                ],
            ]
        )

        imported, errors = import_student_roster(uploaded_file)

        self.assertEqual(imported, 0)
        self.assertTrue(any("different programme" in error for error in errors))
        self.assertFalse(Student.objects.exists())

    def test_template_download_and_form_errors_render(self):
        response = self.client.get(reverse("core:student_import_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("Admission No.", response.content.decode())

        response = self.client.post(
            reverse("core:student_import"),
            {
                "roster_file": SimpleUploadedFile(
                    "students.txt",
                    b"unsupported",
                    content_type="text/plain",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Upload a CSV or Excel .xlsx file.")

    def test_management_command_imports_a_validated_csv_roster(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-COMMAND-1",
                    "IDX-COMMAND-1",
                    "Command Import Student",
                    "Male",
                    "General Science",
                    "Red House",
                    "1A1",
                ],
            ]
        )
        with tempfile.TemporaryDirectory() as directory:
            roster_path = Path(directory, uploaded_file.name)
            roster_path.write_bytes(uploaded_file.read())
            output = io.StringIO()

            call_command("import_roster", str(roster_path), stdout=output)

        self.assertIn("1 student record(s) imported", output.getvalue())
        self.assertTrue(
            Student.objects.filter(admission_no="ADM-COMMAND-1").exists()
        )

    def test_management_command_reports_row_errors_without_partial_imports(self):
        uploaded_file = self._csv_upload(
            [
                [
                    "ADM-COMMAND-GOOD",
                    "IDX-COMMAND-GOOD",
                    "Valid Student",
                    "Male",
                    "General Science",
                    "Red House",
                    "1A1",
                ],
                [
                    "ADM-COMMAND-BAD",
                    "IDX-COMMAND-BAD",
                    "Invalid Student",
                    "Female",
                    "Unknown Programme",
                    "Red House",
                    "1A1",
                ],
            ]
        )
        with tempfile.TemporaryDirectory() as directory:
            roster_path = Path(directory, uploaded_file.name)
            roster_path.write_bytes(uploaded_file.read())

            with self.assertRaisesMessage(CommandError, "CSV, row 3"):
                call_command("import_roster", str(roster_path))

        self.assertFalse(Student.objects.exists())


class TemporaryStudentMediaTests(TransactionTestCase):
    def test_registry_displays_student_photo_as_a_cropped_portrait(self):
        with tempfile.TemporaryDirectory() as media_dir:
            with override_settings(
                MEDIA_ROOT=media_dir,
                PRIVATE_STUDENT_PHOTO_ROOT=media_dir,
            ):
                user = get_user_model().objects.create_user(
                    username="photo-admin",
                    password="password123",
                )
                UserProfile.objects.update_or_create(
                    user=user,
                    defaults={"role": StaffRole.ADMIN},
                )
                programme = Programme.objects.create(name="Arts")
                house = House.objects.create(name="Blue House")
                dormitory = Dormitory.objects.create(
                    house=house,
                    name="Girls Block",
                    gender=Dormitory.GenderBlock.FEMALE,
                )
                image = Image.new("RGB", (100, 140), color="teal")
                image_data = io.BytesIO()
                image.save(image_data, format="PNG")
                student = Student.objects.create(
                    admission_no="ADM-PHOTO",
                    index_number="IDX-PHOTO",
                    full_name="Photo Student",
                    gender=Student.Gender.FEMALE,
                    programme=programme,
                    house=house,
                    dormitory=dormitory,
                    photo=SimpleUploadedFile(
                        "portrait.png",
                        image_data.getvalue(),
                        content_type="image/png",
                    ),
                )
                self.client.force_login(user)

                response = self.client.get(reverse("core:student_registry"))

                self.assertEqual(response.status_code, 200)
                self.assertContains(
                    response,
                    reverse("core:student_photo", args=[student.pk]),
                )
                self.assertContains(response, "object-cover")
                self.assertTrue(Path(media_dir, student.photo.name).exists())
                photo_response = self.client.get(
                    reverse("core:student_photo", args=[student.pk])
                )
                self.assertEqual(photo_response.status_code, 200)
                self.assertEqual(photo_response["X-Content-Type-Options"], "nosniff")
                photo_response.close()
                with self.assertRaises(ValueError):
                    student.photo.url
